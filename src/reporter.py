"""성과 Excel — 매일 KST 10:00 수집, 날짜×언어 4행 스냅샷 누적. 알림 없음."""
from __future__ import annotations

from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook

from .error_log import log_error
from .generator import LANGS
from .publishers import PUBLISHERS

KST = timezone(timedelta(hours=9))
FILE_NAME = "performance.xlsx"
HEADERS = ["date", "lang", "views_impressions", "likes", "comments", "shares", "saves", "collected_at"]


def _path(data_dir: Path) -> Path:
    return data_dir / FILE_NAME


def collect_all(data_dir: Path, mock: bool = False) -> list[dict]:
    data_dir.mkdir(parents=True, exist_ok=True)
    path = _path(data_dir)
    wb = load_workbook(path) if path.exists() else Workbook()
    ws = wb.active
    if ws["A1"].value != "date":
        for col, h in enumerate(HEADERS, start=1):
            ws.cell(row=1, column=col, value=h)

    now = datetime.now(KST)
    today = now.date().isoformat()
    collected = []

    for lang in LANGS:
        total = {"views": 0, "likes": 0, "comments": 0, "shares": 0, "saves": 0}
        if not mock:
            for pub in PUBLISHERS:
                try:
                    m = pub.fetch_metrics(lang)
                except Exception as e:  # noqa: BLE001
                    log_error(data_dir, f"metrics:{pub.name}:{lang}", str(e))
                    m = None
                if m:
                    for k in total:
                        total[k] += m.get(k, 0)
        row = [today, lang, total["views"], total["likes"], total["comments"],
               total["shares"], total["saves"], now.isoformat(timespec="seconds")]
        ws.append(row)
        collected.append(dict(zip(HEADERS, row)))

    wb.save(path)
    return collected
