"""소스 Excel — 내부 기록용 (사이트명/원문 URL 보관, 스냅샷 누적)."""
from __future__ import annotations

from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook

KST = timezone(timedelta(hours=9))
FILE_NAME = "source_log.xlsx"
HEADERS = ["timestamp_kst", "date", "cycle", "source", "post_id", "post_url",
           "title", "views", "comments_used", "status", "note"]


def _path(data_dir: Path) -> Path:
    return data_dir / FILE_NAME


def _open(data_dir: Path):
    path = _path(data_dir)
    if path.exists():
        return load_workbook(path)
    wb = Workbook()
    wb.active.append(HEADERS)
    return wb


def append(data_dir: Path, cycle: str, post, status: str, note: str = "") -> None:
    data_dir.mkdir(parents=True, exist_ok=True)
    wb = _open(data_dir)
    now = datetime.now(KST)
    wb.active.append([
        now.isoformat(timespec="seconds"), now.date().isoformat(), cycle,
        post.source, post.id, post.url, post.title,
        post.views, len(post.comments), status, note,
    ])
    wb.save(_path(data_dir))


def used_post_ids_today(data_dir: Path, cycle: str = "") -> set[str]:
    """오늘 사용한(게시 시도된) 게시물 id 전체 목록.

    PRD: 06:00 게시물이 12:00에도 1위면 다음 순위 사용 → 하루에 한 번이라도
    사용된 게시물은 같은 날 재선정 금지 (사이클 무관, 사이클 재실행 포함).
    published + partial_error 모두 '실제 게시 시도'로 간주 (mock 제외).
    """
    path = _path(data_dir)
    if not path.exists():
        return set()
    today = datetime.now(KST).date().isoformat()
    used = set()
    for row in load_workbook(path).active.iter_rows(min_row=2, values_only=True):
        if row[1] == today and row[9] in ("published", "partial_error"):
            used.add(row[4])
    return used
